#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera automáticamente un informe Word del Caso 2 (Sedimentos) de CIMMA.

Carpeta esperada:
  datos/sedimentos_SD01.csv
  figuras/figura3_color_natural_enunciado.png   (referencia del PDF del taller)
  figuras/figura4_falso_color_enunciado.png     (referencia del PDF del taller)
  imagenes_gee/                                 (opcional: exports .tif/.tiff de GEE)

Para incluir mapas propios de GEE, descarga los GeoTIFF de Google Drive y colócalos en
imagenes_gee/. Nombres reconocidos:
  Natural_2017_B4_B3_B2.tif
  Natural_2025_B4_B3_B2.tif
  FalsoColor_2017_B11_B8_B4.tif
  FalsoColor_2025_B11_B8_B4.tif
  Cambio_NDVI_2025_menos_2017.tif

Ejecución desde la carpeta raíz del paquete:
  python codigo/generar_informe_caso2.py

Requisitos: pandas, matplotlib, python-docx, numpy, rasterio (este último es opcional
si solo se usan las figuras del enunciado y no se procesan GeoTIFF).
Los datos analíticos se transcriben de la tabla incluida en la actividad; no son nuevas
mediciones de campo. Este script no detecta ni mide mercurio en imágenes satelitales.
"""
from __future__ import annotations

from pathlib import Path
import math
import sys
import warnings

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    import rasterio
    from rasterio.enums import Resampling
except Exception:
    rasterio = None
    Resampling = None

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "datos" / "sedimentos_SD01.csv"
FIG_DIR = ROOT / "salidas" / "figuras_informe"
GEE_DIR = ROOT / "imagenes_gee"
PDF_FIG_DIR = ROOT / "figuras"
OUT_DIR = ROOT / "salidas"
OUT_DOCX = OUT_DIR / "Informe_Caso2_Sedimentos_CIMMA.docx"

# Recuento visto en la ejecución compartida por el usuario el 18/10/2026.
# Actualizar estos números si cambia la consulta en Google Earth Engine.
SCENES_2017 = 2
SCENES_2025 = 44

# Guías canadienses para sedimentos de agua dulce (mg/kg de peso seco).
# Son referencias comparativas internacionales, NO ECA peruanos.
REFERENCE = [
    ("Mercurio total", "mercurio_total_mg_kg", 0.170, 0.486, "https://ccme.ca/en/chemical/131"),
    ("Arsénico", "arsenico_mg_kg", 5.9, 17.0, "https://ccme.ca/en/chemical/9"),
    ("Plomo", "plomo_mg_kg", 35.0, 91.3, "https://www.ccme.ca/en/res/lead-canadian-sediment-quality-guidelines-for-the-protection-of-aquatic-life-en.pdf"),
    ("Cadmio", "cadmio_mg_kg", 0.60, 3.50, "https://ccme.ca/en/chemical/20"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, bold=True, color=(255, 255, 255), size=font_size)
        set_cell_shading(table.rows[0].cells[j], "1F4E78")
    # repeat header on multiple pages
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, v in enumerate(row):
            set_cell_text(cells[j], v, size=font_size)
            if i % 2 == 1:
                set_cell_shading(cells[j], "EEF3F8")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def fmt_num(v, digits=2):
    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
        return "—"
    return f"{float(v):,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_pct(v, digits=1):
    return fmt_num(v, digits) + "%"


def save_charts(df: pd.DataFrame) -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}

    charts = [
        ("mercurio", "Mercurio total en sedimento de fondo | SD-01", "mercurio_total_mg_kg", "Mercurio total (mg/kg)"),
        ("metilmercurio", "Metilmercurio en sedimento de fondo | SD-01", "metilmercurio_ug_kg", "Metilmercurio (µg/kg)"),
        ("arsenico", "Arsénico en sedimento de fondo | SD-01", "arsenico_mg_kg", "Arsénico (mg/kg)"),
        ("plomo", "Plomo en sedimento de fondo | SD-01", "plomo_mg_kg", "Plomo (mg/kg)"),
        ("cadmio", "Cadmio en sedimento de fondo | SD-01", "cadmio_mg_kg", "Cadmio (mg/kg)"),
        ("materia_organica", "Materia orgánica del sedimento | SD-01", "materia_organica_pct", "Materia orgánica (%)"),
    ]
    for key, title, col, ylabel in charts:
        fig, ax = plt.subplots(figsize=(8.6, 4.4), constrained_layout=True)
        ax.plot(df["anio"], df[col], marker="o", linewidth=2)
        ax.set_title(title, fontsize=12)
        ax.set_xlabel("Año de muestreo")
        ax.set_ylabel(ylabel)
        ax.set_xticks(df["anio"].astype(int))
        ax.tick_params(axis="x", rotation=35)
        ax.grid(True, alpha=0.28)
        p = FIG_DIR / f"{key}.png"
        fig.savefig(p, dpi=180, bbox_inches="tight")
        plt.close(fig)
        paths[key] = p

    fig, axes = plt.subplots(2, 2, figsize=(11, 7.6), constrained_layout=True)
    series = [
        (axes[0, 0], "mercurio_total_mg_kg", "Mercurio total (mg/kg)"),
        (axes[0, 1], "metilmercurio_ug_kg", "Metilmercurio (µg/kg)"),
        (axes[1, 1], "materia_organica_pct", "Materia orgánica (%)"),
    ]
    for ax, col, title in series:
        ax.plot(df["anio"], df[col], marker="o", linewidth=1.8)
        ax.set_title(title)
        ax.set_xlabel("Año")
        ax.grid(True, alpha=0.25)
        ax.set_xticks(df["anio"].astype(int))
        ax.tick_params(axis="x", rotation=30)
    ax = axes[1, 0]
    for col, label in [("arsenico_mg_kg", "Arsénico"), ("plomo_mg_kg", "Plomo"), ("cadmio_mg_kg", "Cadmio")]:
        ax.plot(df["anio"], df[col] / df[col].iloc[0] * 100, marker="o", label=label)
    ax.set_title("Otros metales (2017 = 100)")
    ax.set_xlabel("Año")
    ax.set_ylabel("Índice relativo (%)")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(True, alpha=0.25)
    ax.set_xticks(df["anio"].astype(int))
    ax.tick_params(axis="x", rotation=30)
    panel = FIG_DIR / "panel_evolucion.png"
    fig.suptitle("Caso 2: evolución temporal de los sedimentos, 2017–2025", fontsize=14)
    fig.savefig(panel, dpi=180, bbox_inches="tight")
    plt.close(fig)
    paths["panel"] = panel

    # Calcula tabla de cambios; columnas de datos se conservan en el CSV fuente.
    rows = []
    metric_defs = [
        ("Mercurio total", "mercurio_total_mg_kg", "mg/kg"),
        ("Metilmercurio", "metilmercurio_ug_kg", "µg/kg"),
        ("Arsénico", "arsenico_mg_kg", "mg/kg"),
        ("Plomo", "plomo_mg_kg", "mg/kg"),
        ("Cadmio", "cadmio_mg_kg", "mg/kg"),
        ("Materia orgánica", "materia_organica_pct", "%"),
    ]
    for name, col, unit in metric_defs:
        start = float(df[col].iloc[0])
        end = float(df[col].iloc[-1])
        rows.append({"parametro": name, "valor_2017": start, "valor_2025": end,
                     "unidad": unit, "cambio_absoluto": end-start,
                     "factor_2025_vs_2017": end/start,
                     "variacion_porcentual": (end-start)/start*100})
    pd.DataFrame(rows).to_csv(ROOT / "datos" / "resumen_cambios_calculado.csv", index=False, float_format="%.4f")
    return paths


def find_file(patterns: list[str]) -> Path | None:
    folders = [GEE_DIR, ROOT]
    for folder in folders:
        if not folder.exists():
            continue
        for pat in patterns:
            hits = list(folder.glob(pat)) if folder == GEE_DIR else list(folder.rglob(pat))
            hits = [p for p in hits if p.is_file() and p.suffix.lower() in {".tif", ".tiff", ".png", ".jpg", ".jpeg"}]
            if hits:
                return hits[0]
    return None


def raster_to_png(path: Path, output: Path, mode: str):
    """Renderiza un GeoTIFF exportado por GEE para insertarlo en Word."""
    output.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        from PIL import Image
        Image.open(path).convert("RGB").save(output)
        return output
    if rasterio is None:
        raise RuntimeError("Necesitas instalar rasterio para convertir GeoTIFF: pip install rasterio")
    with rasterio.open(path) as src:
        max_dim = 1800
        scale = min(1, max_dim / max(src.width, src.height))
        out_w = max(1, int(src.width * scale))
        out_h = max(1, int(src.height * scale))
        if mode == "ndvi":
            arr = src.read(1, out_shape=(out_h, out_w), resampling=Resampling.bilinear, masked=True)
            fig, ax = plt.subplots(figsize=(8, 5.2), constrained_layout=True)
            im = ax.imshow(arr, cmap="RdYlGn", vmin=-0.5, vmax=0.5)
            ax.set_title("Diferencia de NDVI (2025 − 2017)")
            ax.set_axis_off()
            fig.colorbar(im, ax=ax, shrink=0.75, label="ΔNDVI")
            fig.savefig(output, dpi=170, bbox_inches="tight")
            plt.close(fig)
        else:
            if src.count >= 3:
                arr = src.read([1, 2, 3], out_shape=(3, out_h, out_w), resampling=Resampling.nearest, masked=True)
                arr = np.ma.filled(arr, 0)
                arr = np.moveaxis(arr, 0, -1)
                if np.nanmax(arr) <= 1.0:
                    arr = arr * 255.0
                arr = np.clip(arr, 0, 255).astype(np.uint8)
                from PIL import Image
                Image.fromarray(arr, mode="RGB").save(output)
            else:
                arr = src.read(1, out_shape=(out_h, out_w), resampling=Resampling.bilinear, masked=True)
                fig, ax = plt.subplots(figsize=(8, 5), constrained_layout=True)
                ax.imshow(arr, cmap="gray")
                ax.set_axis_off()
                fig.savefig(output, dpi=170, bbox_inches="tight")
                plt.close(fig)
    return output


def make_gee_figures() -> dict[str, Path]:
    out = {}
    definitions = {
        "natural_2017": (["Natural_2017_B4_B3_B2.tif", "Natural_2017_B4_B3_B2.tiff", "*Natural*2017*B4*B3*B2*"], "rgb"),
        "natural_2025": (["Natural_2025_B4_B3_B2.tif", "Natural_2025_B4_B3_B2.tiff", "*Natural*2025*B4*B3*B2*"], "rgb"),
        "false_2017": (["FalsoColor_2017_B11_B8_B4.tif", "FalsoColor_2017_B11_B8_B4.tiff", "*FalsoColor*2017*"], "rgb"),
        "false_2025": (["FalsoColor_2025_B11_B8_B4.tif", "FalsoColor_2025_B11_B8_B4.tiff", "*FalsoColor*2025*"], "rgb"),
        "delta_ndvi": (["Cambio_NDVI_2025_menos_2017.tif", "Cambio_NDVI_2025_menos_2017.tiff", "*Cambio*NDVI*2025*2017*"], "ndvi"),
    }
    for key, (patterns, mode) in definitions.items():
        source = find_file(patterns)
        if source:
            target = FIG_DIR / f"gee_{key}.png"
            try:
                out[key] = raster_to_png(source, target, mode)
                print(f"Mapa GEE detectado: {source.name}")
            except Exception as exc:
                warnings.warn(f"No se pudo convertir {source}: {exc}")
    return out


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)


def add_figure(doc, path: Path | None, caption: str, width=6.2):
    if path and path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)
    else:
        doc.add_paragraph(f"[Figura pendiente: {caption}]")


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def build_report():
    if not DATA.exists():
        raise FileNotFoundError(f"No encuentro el CSV de datos: {DATA}")
    df = pd.read_csv(DATA)
    required = {"anio", "mercurio_total_mg_kg", "metilmercurio_ug_kg", "arsenico_mg_kg", "plomo_mg_kg", "cadmio_mg_kg", "materia_organica_pct"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Al CSV le faltan columnas: {', '.join(sorted(missing))}")

    charts = save_charts(df)
    gee = make_gee_figures()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for style_name, size, color in [("Title", 24, "17365D"), ("Heading 1", 16, "17365D"), ("Heading 2", 12, "1F4E78")]:
        st = styles[style_name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    # Header / footer
    header = sec.header.paragraphs[0]
    header.text = "CIMMA | Eje de Teledetección | Caso 2: Sedimentos"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(100, 100, 100)
    footer = sec.footer.paragraphs[0]
    footer.text = "Borrador técnico reproducible — verificar área GEE y base de concentración antes de versión final"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(100, 100, 100)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EVALUACIÓN MULTITEMPORAL DE SEDIMENTOS")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Minería aurífera aluvial en Huepetuhe, Madre de Dios | 2017–2025")
    r.font.size = Pt(15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Informe automatizado de avance\nCírculo de Minería y Medio Ambiente (CIMMA)\nRol de análisis: Dirección de Evaluación Ambiental (DEAM) – OEFA").font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DOCUMENTO DE TRABAJO — requiere validación del área satelital y revisión de fuentes")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(166, 77, 0)
    doc.add_page_break()

    doc.add_heading("Resumen ejecutivo", level=1)
    hg_start = float(df["mercurio_total_mg_kg"].iloc[0]); hg_end = float(df["mercurio_total_mg_kg"].iloc[-1])
    mm_start = float(df["metilmercurio_ug_kg"].iloc[0]); mm_end = float(df["metilmercurio_ug_kg"].iloc[-1])
    om_start = float(df["materia_organica_pct"].iloc[0]); om_end = float(df["materia_organica_pct"].iloc[-1])
    doc.add_paragraph(
        f"La tabla de monitoreo entregada en el taller muestra que en SD-01 el mercurio total pasó de {fmt_num(hg_start)} a {fmt_num(hg_end)} mg/kg entre 2017 y 2025 (factor {fmt_num(hg_end/hg_start, 2)}; incremento de {fmt_pct((hg_end-hg_start)/hg_start*100)}). "
        f"El metilmercurio aumentó de {fmt_num(mm_start,1)} a {fmt_num(mm_end,1)} µg/kg (factor {fmt_num(mm_end/mm_start,2)}), mientras que la materia orgánica bajó de {fmt_num(om_start,1)} % a {fmt_num(om_end,1)} % (reducción de {fmt_pct((om_start-om_end)/om_start*100)}). "
        "Las imágenes suministradas en el taller muestran expansión de superficies expuestas y fragmentación de la cobertura boscosa. La relación espacial y temporal es consistente con una mayor presión minera, pero los datos de un único punto y las imágenes por sí solas no demuestran causalidad."
    )
    doc.add_paragraph(
        f"Estado de teledetección: en la ejecución reportada del script de Earth Engine se registraron {SCENES_2017} escenas para 2017 y {SCENES_2025} para 2025; además, el área de análisis se definió de forma aproximada porque el enunciado no incluye coordenadas exactas de SD-01. Por tanto, los mapas propios de GEE y la cuantificación de superficies deben tratarse como preliminares hasta validar estos dos aspectos."
    )

    doc.add_heading("1. Objetivo y alcance", level=1)
    doc.add_paragraph(
        "Evaluar la evolución 2017–2025 de los sedimentos de fondo del río Huepetuhe integrando (i) interpretación multitemporal de imágenes satelitales, (ii) tendencias de las concentraciones de mercurio total, metilmercurio, arsénico, plomo y cadmio, y (iii) variación de la materia orgánica. El análisis utiliza los datos didácticos proporcionados por el taller y no sustituye una evaluación ambiental oficial ni nuevos resultados de laboratorio."
    )

    doc.add_heading("2. Área de estudio y contexto", level=1)
    doc.add_paragraph(
        "Huepetuhe se ubica en la provincia de Manu, región Madre de Dios. El taller describe minería aurífera aluvial activa, con remoción de vegetación y suelo, generación de pozas, depósitos de material y posibles aportes de partículas a los cauces. Estos procesos pueden modificar la geomorfología y favorecer el transporte y la acumulación de contaminantes en sedimentos."
    )

    doc.add_heading("3. Datos y metodología", level=1)
    doc.add_heading("3.1 Datos de sedimentos", level=2)
    doc.add_paragraph(
        "Se utilizaron los valores anuales 2017–2025 de la tabla del taller para el punto SD-01. El conjunto cuenta con una campaña por año. La transcripción está guardada en datos/sedimentos_SD01.csv dentro del paquete. Antes de comparar con guías de sedimentos, debe confirmarse que los resultados estén expresados en peso seco."
    )
    doc.add_heading("3.2 Procesamiento satelital en Google Earth Engine", level=2)
    doc.add_paragraph(
        "El script usa COPERNICUS/S2_SR_HARMONIZED y filtra la estación seca (1 de mayo a 1 de octubre) de 2017 y 2025. Aplica una máscara basada en la clasificación SCL, calcula composiciones medianas y genera color natural (B4-B3-B2), falso color SWIR-NIR-rojo (B11-B8-B4), NDVI y la diferencia NDVI(2025) − NDVI(2017). El área usada en la primera ejecución es un buffer aproximado de 6 km; no corresponde a un polígono de levantamiento validado."
    )
    doc.add_paragraph(
        f"La consola observada mostró {SCENES_2017} escenas en 2017 y {SCENES_2025} en 2025. La desigualdad del número de escenas puede afectar la comparabilidad de las composiciones; se recomienda ampliar/validar la búsqueda temporal y revisar la fracción de píxeles válidos de cada compuesto antes de extraer conclusiones finales."
    )

    doc.add_heading("4. Resultados de la comparación satelital", level=1)
    if "natural_2017" in gee and "natural_2025" in gee:
        fig, ax = plt.subplots(1, 2, figsize=(10, 4.8), constrained_layout=True)
        from PIL import Image
        for a, key, title in [(ax[0], "natural_2017", "Color natural — 2017"), (ax[1], "natural_2025", "Color natural — 2025")]:
            a.imshow(Image.open(gee[key])); a.set_title(title); a.axis("off")
        composite = FIG_DIR / "gee_comparacion_natural.png"
        fig.savefig(composite, dpi=180, bbox_inches="tight"); plt.close(fig)
        add_figure(doc, composite, "Figura 1. Composiciones propias de GEE. Validar que el AOI coincide con la zona SD-01.")
    else:
        natural_source = PDF_FIG_DIR / "figura3_color_natural_enunciado.png"
        add_figure(doc, natural_source if natural_source.exists() else None,
                   "Figura 1. Comparación de color natural reproducida del enunciado (p. 5); no corresponde a una exportación propia de GEE.")

    if "false_2017" in gee and "false_2025" in gee:
        fig, ax = plt.subplots(1, 2, figsize=(10, 4.8), constrained_layout=True)
        from PIL import Image
        for a, key, title in [(ax[0], "false_2017", "Falso color — 2017"), (ax[1], "false_2025", "Falso color — 2025")]:
            a.imshow(Image.open(gee[key])); a.set_title(title); a.axis("off")
        composite = FIG_DIR / "gee_comparacion_falsocolor.png"
        fig.savefig(composite, dpi=180, bbox_inches="tight"); plt.close(fig)
        add_figure(doc, composite, "Figura 2. Composiciones SWIR-NIR-rojo propias de GEE. La interpretación se limita a cambios espectrales.")
    else:
        false_source = PDF_FIG_DIR / "figura4_falso_color_enunciado.png"
        add_figure(doc, false_source if false_source.exists() else None,
                   "Figura 2. Combinación B11-B8-B4 reproducida del enunciado (p. 5); no corresponde a una exportación propia de GEE.")

    if "delta_ndvi" in gee:
        add_figure(doc, gee["delta_ndvi"], "Figura 3. Diferencia de NDVI (2025 − 2017) exportada de GEE. Los valores negativos indican descenso del índice, no contaminación química.")
    else:
        doc.add_paragraph(
            "La figura de diferencia NDVI de la ejecución propia de GEE todavía debe incorporarse. Al descargar el archivo Cambio_NDVI_2025_menos_2017.tif desde Drive, colóquelo en imagenes_gee/ y vuelva a ejecutar este script."
        )

    doc.add_heading("4.1 Interpretación visual (pregunta a)", level=2)
    doc.add_paragraph(
        "En las imágenes del taller, la comparación 2017–2025 muestra más superficies claras y discontinuas alrededor del cauce en 2025, compatibles visualmente con terrenos removidos, pozas y depósitos expuestos, junto con una mayor fragmentación del bosque. Se clasifica cualitativamente el cambio de superficie intervenida como ALTO. Esta clasificación no es un cálculo de hectáreas: para estimar el área es necesario delimitar clases en un SIG, usar imágenes comparables y documentar la exactitud de la clasificación."
    )
    doc.add_paragraph(
        "La combinación de falso color aporta contraste entre la vegetación y las superficies con distinta respuesta espectral en SWIR/NIR/rojo. Sin embargo, las respuestas espectrales no permiten identificar por sí solas el mercurio ni otros metales en sedimentos. La teledetección sirve para localizar zonas de transformación y orientar el muestreo de campo."
    )

    doc.add_heading("5. Resultados del monitoreo de sedimentos", level=1)
    doc.add_paragraph("La siguiente tabla compara el primer y último año de la serie; los porcentajes se calculan a partir de los valores del taller.")
    results = []
    metrics = [
        ("Mercurio total", "mercurio_total_mg_kg", "mg/kg"),
        ("Metilmercurio", "metilmercurio_ug_kg", "µg/kg"),
        ("Arsénico", "arsenico_mg_kg", "mg/kg"),
        ("Plomo", "plomo_mg_kg", "mg/kg"),
        ("Cadmio", "cadmio_mg_kg", "mg/kg"),
        ("Materia orgánica", "materia_organica_pct", "%"),
    ]
    for label, col, unit in metrics:
        a, b = float(df[col].iloc[0]), float(df[col].iloc[-1])
        results.append([label, f"{fmt_num(a)} {unit}", f"{fmt_num(b)} {unit}", fmt_num(b/a, 2) + "×", fmt_pct((b-a)/a*100)])
    add_table(doc, ["Parámetro", "2017", "2025", "Factor final/inicial", "Variación"], results, font_size=8.2)

    chart_specs = [
        ("mercurio", "Figura 4. Evolución anual del mercurio total en el punto SD-01."),
        ("metilmercurio", "Figura 5. Evolución anual del metilmercurio en el punto SD-01."),
        ("arsenico", "Figura 6. Evolución anual del arsénico en el punto SD-01."),
        ("plomo", "Figura 7. Evolución anual del plomo en el punto SD-01."),
        ("cadmio", "Figura 8. Evolución anual del cadmio en el punto SD-01."),
        ("materia_organica", "Figura 9. Evolución anual de la materia orgánica del sedimento en el punto SD-01."),
    ]
    for key, caption in chart_specs:
        add_figure(doc, charts[key], caption, width=6.0)

    doc.add_heading("5.1 Pregunta b: incremento del mercurio total", level=2)
    doc.add_paragraph(
        f"El mercurio total aumentó de {fmt_num(hg_start)} a {fmt_num(hg_end)} mg/kg. El factor de cambio es C2025/C2017 = {fmt_num(hg_end/hg_start, 2)} y el incremento porcentual es ((C2025 − C2017)/C2017) × 100 = {fmt_pct((hg_end-hg_start)/hg_start*100)}. La tendencia temporal coincide con la expansión de áreas intervenidas descrita en las imágenes. La asociación es compatible con mayor transporte/acumulación de materiales, pero el diseño de un único punto anual no prueba por sí mismo que la minería sea la única fuente."
    )

    doc.add_heading("5.2 Pregunta c: metilmercurio y condiciones de las pozas", level=2)
    doc.add_paragraph(
        f"El metilmercurio pasó de {fmt_num(mm_start,1)} a {fmt_num(mm_end,1)} µg/kg (factor {fmt_num(mm_end/mm_start,2)}). Las pozas con circulación limitada pueden formar zonas pobres en oxígeno en sedimentos, donde microorganismos específicos transforman parte del mercurio inorgánico en metilmercurio. El proceso depende además de disponibilidad de materia orgánica, azufre, temperatura, pH y condiciones redox; no se puede inferir la tasa de metilación solo a partir de la concentración observada. El metilmercurio es especialmente preocupante por su incorporación a la cadena trófica y su bioacumulación/biomagnificación."
    )
    # Distinguish period change from annual intervals
    hg_annual = df["mercurio_total_mg_kg"].pct_change() * 100
    mm_annual = df["metilmercurio_ug_kg"].pct_change() * 100
    years_more = []
    for i in range(1, len(df)):
        if mm_annual.iloc[i] > hg_annual.iloc[i]:
            years_more.append(f"{int(df['anio'].iloc[i-1])}–{int(df['anio'].iloc[i])}")
    doc.add_paragraph(
        "Matiz cuantitativo: el factor acumulado del mercurio total (" + fmt_num(hg_end/hg_start,2) + "×) es mayor que el del metilmercurio (" + fmt_num(mm_end/mm_start,2) + "×). El metilmercurio crece proporcionalmente más que el mercurio total en algunos intervalos anuales" + (" (" + ", ".join(years_more) + ")." if years_more else ", aunque no se identificaron intervalos en esta tabla.")
    )

    doc.add_heading("5.3 Pregunta d: materia orgánica y material estéril", level=2)
    doc.add_paragraph(
        f"La materia orgánica bajó de {fmt_num(om_start,1)} % a {fmt_num(om_end,1)} %, una reducción de {fmt_pct((om_start-om_end)/om_start*100)}. No es necesariamente contradictorio que los metales aumenten al mismo tiempo: la pérdida de vegetación y de la capa superficial puede reducir el aporte de hojarasca y materia orgánica, mientras que los relaves y el material estéril de origen mineral pueden diluir el porcentaje de materia orgánica y aportar partículas que contienen metales. Estos mecanismos son plausibles, pero requieren datos de granulometría, carbono orgánico total, mineralogía y geoquímica para ser verificados en el sitio."
    )

    doc.add_heading("6. Priorización de metales y comparación referencial", level=1)
    doc.add_paragraph(
        "El enunciado solicita revisar ECA u otro referente. La tabla siguiente compara los datos de 2025 con guías de sedimentos de agua dulce del Canadian Council of Ministers of the Environment (CCME). Son referencias internacionales ecotoxicológicas, no límites legales peruanos. La comparación solo es válida si los datos del taller están expresados en peso seco y sus métodos analíticos son comparables."
    )
    riskrows = []
    reference_precision = {"mercurio_total_mg_kg": 3, "arsenico_mg_kg": 1, "plomo_mg_kg": 1, "cadmio_mg_kg": 2}
    for label, col, isqg, pel, url in REFERENCE:
        value = float(df[col].iloc[-1])
        if value < isqg:
            status = "Por debajo de ISQG"
        elif value < pel:
            status = "Entre ISQG y PEL"
        else:
            status = "Igual/superior a PEL"
        p_digits = reference_precision[col]
        riskrows.append([label, fmt_num(value), fmt_num(isqg, p_digits), fmt_num(pel, p_digits), status])
    add_table(doc, ["Metal", "Resultado 2025 (mg/kg)", "ISQG (mg/kg)", "PEL (mg/kg)", "Lectura referencial"], riskrows, font_size=7.8)
    doc.add_paragraph(
        "Con los supuestos anteriores, mercurio total es la prioridad de seguimiento porque 3,80 mg/kg supera la guía PEL de 0,486 mg/kg. Arsénico, plomo y cadmio se ubican entre sus respectivos valores ISQG y PEL. Esta clasificación es preliminar y no equivale a declarar incumplimiento legal ni a calcular el riesgo ecológico completo. No se asigna aquí un umbral al metilmercurio utilizando los valores del mercurio total."
    )

    doc.add_heading("7. Recomendaciones de evaluación y seguimiento", level=1)
    for text in [
        "Validar el área de interés con coordenadas oficiales del punto SD-01 y ajustar el polígono de análisis al cauce y las superficies mineras evaluadas.",
        "Mejorar la comparabilidad satelital: usar composiciones de estación equivalente, revisar la cantidad de píxeles válidos después de enmascarar nubes y sombras, y documentar la cobertura de cada año. La cuenta de escenas brutas no sustituye esta validación.",
        "Cuantificar el cambio de cobertura mediante clasificación o digitalización de polígonos y reportar hectáreas solo después de la verificación visual y, cuando sea posible, validación de campo.",
        "Mantener campañas de sedimentos con control de calidad, replicados y puntos de referencia aguas arriba y aguas abajo; confirmar peso seco, granulometría y materia orgánica.",
        "Priorizar mercurio total y metilmercurio; incluir arsénico, plomo y cadmio. Considerar análisis de tejidos de peces cuando el diseño de evaluación y las competencias de la entidad lo justifiquen.",
        "Usar la teledetección para detectar zonas prioritarias y orientar el muestreo; no interpretar los índices espectrales como mediciones directas de concentración de metales."
    ]:
        add_bullet(doc, text)

    doc.add_heading("8. Conclusiones", level=1)
    conclusions = [
        "Las imágenes del taller sugieren un aumento alto de la superficie intervenida y una mayor fragmentación del bosque entre 2017 y 2025; no se calculó un área exacta en hectáreas.",
        f"En SD-01, el mercurio total aumentó {fmt_num(hg_end/hg_start,2)} veces y el metilmercurio {fmt_num(mm_end/mm_start,2)} veces en el periodo 2017–2025.",
        f"Arsénico, plomo y cadmio también aumentaron, mientras que la materia orgánica disminuyó {fmt_pct((om_start-om_end)/om_start*100)}.",
        "Los patrones son consistentes con una presión ambiental creciente, pero la causalidad y la extensión espacial del riesgo requieren más puntos de muestreo, validación del área satelital y análisis de laboratorio complementario."
    ]
    for c in conclusions:
        add_bullet(doc, c)

    doc.add_heading("Referencias y trazabilidad", level=1)
    refs = [
        "CIMMA – Eje de Teledetección. Actividad integrada: Teledetección y monitoreo de campo. Evolución ambiental de la minería aluvial en Huepetuhe, Madre de Dios (2017–2025). Documento de taller proporcionado para el Caso 2, pp. 5–6.",
        "Google Earth Engine Data Catalog. Harmonized Sentinel-2 MSI: MultiSpectral Instrument, Level-2A. https://developers.google.com/earth-engine/datasets/catalog/COPERNICUS_S2_SR_HARMONIZED",
        "CCME. Sediment Quality Guidelines for the Protection of Aquatic Life — Mercury (freshwater). https://ccme.ca/en/chemical/131",
        "CCME. Sediment Quality Guidelines for the Protection of Aquatic Life — Arsenic (freshwater). https://ccme.ca/en/chemical/9",
        "CCME. Canadian Sediment Quality Guidelines for the Protection of Aquatic Life — Lead. https://www.ccme.ca/en/res/lead-canadian-sediment-quality-guidelines-for-the-protection-of-aquatic-life-en.pdf",
        "CCME. Sediment Quality Guidelines for the Protection of Aquatic Life — Cadmium (freshwater). https://ccme.ca/en/chemical/20",
        "Ministerio del Ambiente del Perú. Decreto Supremo N.° 018-2025-MINAM, que aprueba el Protocolo Nacional para el Monitoreo de la Calidad Ambiental de Sedimentos. https://sinia.minam.gob.pe/sites/default/files/archivos/public/docs/DS%20018-2025-minam.pdf",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    doc.add_heading("Anexo A. Archivos de entrada y regeneración", level=1)
    doc.add_paragraph(
        "Datos de campo: datos/sedimentos_SD01.csv. Código satelital: codigo/Caso2_Huepetuhe_Sentinel2_GEE.js. Para insertar salidas de GEE, descarga los cinco GeoTIFF de la carpeta Caso2_CIMMA en Google Drive y cópialos en imagenes_gee/. Ejecuta de nuevo python codigo/generar_informe_caso2.py. El informe y las figuras se guardarán en salidas/."
    )

    doc.core_properties.title = "Caso 2 CIMMA — Sedimentos de Huepetuhe 2017–2025"
    doc.core_properties.subject = "Informe técnico reproducible con datos del taller y teledetección Sentinel-2"
    doc.core_properties.author = "CIMMA — Borrador preparado para revisión"
    doc.save(OUT_DOCX)

    print("Informe generado:", OUT_DOCX)
    print("Figuras generadas:", FIG_DIR)
    print("Exportaciones GEE incorporadas:", ", ".join(gee.keys()) if gee else "ninguna; se conservaron figuras de referencia del taller")
    print("Nota: valida área de interés y comparabilidad del compuesto 2017 antes de usarlo como resultado definitivo.")
    return OUT_DOCX


if __name__ == "__main__":
    try:
        build_report()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
